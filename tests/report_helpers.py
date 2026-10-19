"""Helpers of the report tests of wave 2: calls with repeated parameters, downloads, reading back
workbooks and PDFs, and the datasets of the expenses and payments reports.

``reports_helpers`` inserts the billing rows the older reports read; this module builds on it.
"""

import io
import re
from dataclasses import dataclass
from datetime import date
from typing import Any

from httpx import AsyncClient, Response
from openpyxl import load_workbook

from app.db.models import Client, Job, Payment
from tests.domain_factories import (
    activity_code_id,
    make_client,
    make_job,
    make_time_entry,
    practice,
    service_id,
    staff,
)
from tests.expense_helpers import make_claim, make_line
from tests.factories import Actor
from tests.reports_helpers import (
    LineSpec,
    MadeInvoice,
    allocate,
    make_invoice,
    payment,
    report_url,
    update_row,
    write_off,
)

SEPTEMBER = {"date_from": date(2026, 9, 1), "date_to": date(2026, 9, 30)}
AUTUMN = {"date_from": date(2026, 9, 1), "date_to": date(2026, 10, 31)}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _pairs(params: dict[str, Any]) -> list[tuple[str, str]]:
    """Query parameters; a list repeats its parameter, ``None`` leaves it out."""
    pairs: list[tuple[str, str]] = []
    for key, value in params.items():
        if value is None:
            continue
        for item in value if isinstance(value, list | tuple) else [value]:
            pairs.append((key, str(item)))
    return pairs


async def call(client: AsyncClient, actor: Actor, org: dict, name: str, **params: Any) -> Response:
    return await client.get(report_url(org, name), headers=actor.headers, params=_pairs(params))


async def get_data(
    client: AsyncClient, actor: Actor, org: dict, name: str, **params: Any
) -> dict[str, Any]:
    """The report's ``data`` ({rows, totals, filters, page}); asserts 200."""
    response = await call(client, actor, org, name, **params)
    assert response.status_code == 200, response.text
    data: dict[str, Any] = response.json()["data"]
    return data


async def download(
    client: AsyncClient, actor: Actor, org: dict, name: str, fmt: str, **params: Any
) -> Response:
    return await call(client, actor, org, name, format=fmt, **params)


def _sheet_rows(content: bytes) -> list[list[Any]]:
    """Every row of a downloaded workbook's sheet, title and filter notes included."""
    sheet = load_workbook(io.BytesIO(content), data_only=True).active
    assert sheet is not None
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def xlsx_table(content: bytes, first_header: str) -> list[list[Any]]:
    """A downloaded workbook's table: the header row, the data rows and the Total row (the title
    and the filter notes above the header are left out)."""
    rows = _sheet_rows(content)
    start = next(i for i, row in enumerate(rows) if row and row[0] == first_header)
    return rows[start:]


def xlsx_lines_above(content: bytes, first_header: str) -> list[str]:
    """The title and filter notes above the table of a downloaded workbook."""
    rows = _sheet_rows(content)
    start = next(i for i, row in enumerate(rows) if row and row[0] == first_header)
    return [str(row[0]) for row in rows[:start] if row and row[0] is not None]


def pdf_text(content: bytes) -> str:
    """The PDF's bytes as text (the pages are not compressed, so the text can be searched)."""
    return content.decode("latin-1")


def pdf_page_count(content: bytes) -> int:
    return len(re.findall(rb"/Type /Page\b", content))


async def make_payment(
    org: dict, client_row: Client, amount: int, *, method: str = "bank_transfer", **fields: Any
) -> Payment:
    """``reports_helpers.payment`` with a payment method."""
    made = await payment(org, client_row, amount, **fields)
    if method != "bank_transfer":
        await update_row(Payment, made.id, method=method)
        made.method = method
    return made


# --- Expenses ------------------------------------------------------------------------------


@dataclass
class ExpenseData:
    owner: Actor
    org: dict
    ann: Actor
    mia: Actor
    acme: Client
    beta: Client
    job: Job


async def expense_dataset(client: AsyncClient) -> ExpenseData:
    """Five lines in three claims (Sept 2026 and early Oct):

    Ann's approved claim: travel 43.00 to Acme's job (billable, recharge 47.30), postage 5.00 to
    Acme (billable, recharge 5.50, already billed) and subsistence 10.00 + 2.00 VAT (overhead,
    paid by the practice). Ann's draft: travel 20.00 for Beta. Mia's submitted claim: mileage
    18.00 for Beta (billable, recharge 19.80).
    """
    owner, org = await practice(client, full_name="Olivia Owner")
    ann = await staff(org, "staff", full_name="Ann Adams")
    mia = await staff(org, "manager", full_name="Mia Manager")
    acme = await make_client(org, "Acme Ltd", code="ACME")
    beta = await make_client(org, "Beta LLP", code="BETA")
    job = await make_job(org, acme, "Accounts 2026", number="J00001")

    approved = await make_claim(org, ann, status="approved", decided_by=mia)
    await make_line(
        org, approved, code="TRAVEL", incurred_on=date(2026, 9, 7), net=4_300, client=acme,
        job=job, is_billable=True, recharge_amount=4_730,
    )  # fmt: skip
    await make_line(
        org, approved, code="POSTAGE", incurred_on=date(2026, 9, 8), net=500, client=acme,
        is_billable=True, recharge_amount=550, billing_status="billed",
    )  # fmt: skip
    await make_line(
        org, approved, code="SUBSISTENCE", incurred_on=date(2026, 9, 9), net=1_000, vat=200,
        payer="practice",
    )  # fmt: skip
    draft = await make_claim(org, ann, status="draft")
    await make_line(
        org, draft, code="TRAVEL", incurred_on=date(2026, 9, 20), net=2_000, client=beta
    )
    submitted = await make_claim(org, mia, status="submitted")
    await make_line(
        org, submitted, code="MILEAGE", incurred_on=date(2026, 10, 2), net=1_800, client=beta,
        is_billable=True, recharge_amount=1_980,
    )  # fmt: skip
    return ExpenseData(owner, org, ann, mia, acme, beta, job)


# --- Payments ------------------------------------------------------------------------------


@dataclass
class PaymentData:
    owner: Actor
    org: dict
    partner: Actor
    acme: Client
    beta: Client
    invoice_acme: MadeInvoice
    invoice_beta: MadeInvoice


async def payment_dataset(client: AsyncClient) -> PaymentData:
    """Four payments: Acme 100.00 by transfer on 5 Sept (60.00 allocated), Beta 50.00 by card on
    20 Sept (all allocated), Acme 25.00 cash on 3 Oct (on account) and a voided Acme cheque of
    10.00 on 15 Sept. Acme's partner is Pat."""
    owner, org = await practice(client, full_name="Olivia Owner")
    partner = await staff(org, "partner", full_name="Pat Partner")
    acme = await make_client(org, "Acme Ltd", code="ACME", partner_user_id=partner.id)
    beta = await make_client(org, "Beta LLP", code="BETA")
    invoice_acme = await make_invoice(
        org, acme, [LineSpec(10_000, vat_bp=0)], issue_date=date(2026, 8, 20)
    )
    invoice_beta = await make_invoice(
        org, beta, [LineSpec(5_000, vat_bp=0)], issue_date=date(2026, 8, 25)
    )
    await make_payment(
        org, acme, 10_000, received_on=date(2026, 9, 5), allocations=[(invoice_acme, 6_000)]
    )
    await make_payment(
        org, beta, 5_000, method="card", received_on=date(2026, 9, 20),
        allocations=[(invoice_beta, 5_000)],
    )  # fmt: skip
    await make_payment(org, acme, 2_500, method="cash", received_on=date(2026, 10, 3))
    await make_payment(
        org, acme, 1_000, method="cheque", received_on=date(2026, 9, 15), voided=True
    )
    return PaymentData(owner, org, partner, acme, beta, invoice_acme, invoice_beta)


# --- Time and billing for the paging and drill-down tests ----------------------------------


@dataclass
class WideData:
    owner: Actor
    org: dict
    sam: Actor
    pat: Actor
    alpha: Client
    beta: Client
    gamma: Client
    a1: Job
    b1: Job
    g1: Job
    acc_id: Any
    accs_id: Any


async def wide_dataset(client: AsyncClient) -> WideData:
    """Three clients with a job each, time by two people in September, an invoice per client,
    write-offs and a payment: enough rows for every report to be paged.

    Alpha (partner Pat, job J00001 on service ACCS): 60 min by Olivia on 7 Sept (activity ACC,
    billed 120.00 on an invoice of 10 Sept) and 30 min by Sam on 8 Sept (written off 20 Sept).
    Beta (J00002): 45 min by Olivia on 14 Sept, invoiced 90.00 + 5% VAT on 15 Sept. Gamma
    (J00003): 60 min by Sam on 15 Sept, invoiced 120.00 on 16 Sept.
    """
    owner, org = await practice(client, full_name="Olivia Owner")
    sam = await staff(org, "staff", full_name="Sam Staff")
    pat = await staff(org, "partner", full_name="Pat Partner")
    alpha = await make_client(org, "Alpha Ltd", code="ALPHA", partner_user_id=pat.id)
    beta = await make_client(org, "Beta LLP", code="BETA")
    gamma = await make_client(org, "Gamma Co", code="GAMMA")
    accs = await service_id(org, "ACCS")
    acc = await activity_code_id(org, "ACC")
    a1 = await make_job(org, alpha, "Accounts 2026", number="J00001", service_id=accs)
    b1 = await make_job(org, beta, "VAT Q3", number="J00002")
    g1 = await make_job(org, gamma, "Payroll", number="J00003")
    t1 = await make_time_entry(
        org, owner, client=alpha, job=a1, entry_date=date(2026, 9, 7), activity_code_id=acc
    )
    t2 = await make_time_entry(
        org, sam, client=alpha, job=a1, entry_date=date(2026, 9, 8), minutes=30
    )
    t3 = await make_time_entry(
        org, owner, client=beta, job=b1, entry_date=date(2026, 9, 14), minutes=45
    )
    t4 = await make_time_entry(org, sam, client=gamma, job=g1, entry_date=date(2026, 9, 15))
    i_alpha = await make_invoice(
        org, alpha, [LineSpec(12_000, job=a1)], issue_date=date(2026, 9, 10)
    )
    i_beta = await make_invoice(
        org, beta, [LineSpec(9_000, vat_bp=500, job=b1)], issue_date=date(2026, 9, 15)
    )
    i_gamma = await make_invoice(
        org, gamma, [LineSpec(12_000, vat_bp=0, job=g1)], issue_date=date(2026, 9, 16)
    )
    await allocate(org, i_alpha, t1)
    await allocate(org, i_beta, t3)
    await allocate(org, i_gamma, t4)
    await write_off(org, alpha, [t2], written_off_on=date(2026, 9, 20), job=a1)
    await payment(org, alpha, 5_000, received_on=date(2026, 9, 30), allocations=[(i_alpha, 5_000)])
    return WideData(owner, org, sam, pat, alpha, beta, gamma, a1, b1, g1, acc, accs)
