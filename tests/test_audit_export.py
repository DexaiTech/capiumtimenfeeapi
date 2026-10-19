"""The audit-log export: CSV and Excel, the filters, the row cap, the formula protection, who may
export, and that the export is itself audited."""

import csv
import io
import json
import uuid
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.services import audit_export
from app.services.xlsx_export import XLSX_MEDIA_TYPE
from app.utils.practice_dates import today_in
from tests.audit_helpers import audit_url, entry_rows, export, insert_entry
from tests.clients_helpers import org_url
from tests.domain_factories import practice, staff
from tests.factories import add_member, create_user, error
from tests.foundation_helpers import add_custom_role

HEADER = [
    "Time (UTC)",
    "User",
    "Email",
    "Action",
    "Event",
    "Summary",
    "Record type",
    "Record id",
    "Changes",
    "IP address",
    "Device",
    "Request id",
    "Details (JSON)",
]
BASE = datetime(2026, 3, 10, 12, 0, tzinfo=UTC)
# The practice's own entries (its creation) are dated today: the window leaves them out.
WINDOW = {
    "from": (BASE - timedelta(minutes=1)).isoformat(),
    "to": (BASE + timedelta(minutes=60)).isoformat(),
}


def at(minutes: int) -> datetime:
    return BASE + timedelta(minutes=minutes)


def csv_table(response) -> list[dict[str, str]]:
    assert response.status_code == 200, response.text
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"), newline="")))
    header, *body = rows
    assert header == HEADER
    return [dict(zip(header, row, strict=True)) for row in body]


def xlsx_sheet(response) -> tuple[list[tuple], list[dict]]:
    """The rows above the table (title and notes) and the table's rows."""
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["Audit log"]
    rows = list(workbook["Audit log"].iter_rows(values_only=True))
    index = next(i for i, row in enumerate(rows) if row[0] == "Time (UTC)")
    assert list(rows[index]) == HEADER
    return rows[:index], [dict(zip(HEADER, row, strict=True)) for row in rows[index + 1 :]]


async def seed(org, owner, other=None):
    """Three entries at fixed times: an invoice issued, a client changed and a payment."""
    other = other or owner
    invoice, client_id = uuid.uuid4(), uuid.uuid4()
    await insert_entry(
        org,
        "invoice.issued",
        actor=owner,
        target_type="invoice",
        target_id=invoice,
        created_at=at(0),
        meta={"number": "INV-00012", "client_name": "Babbage Ltd"},
    )
    await insert_entry(
        org,
        "client.updated",
        actor=other,
        target_type="client",
        target_id=client_id,
        created_at=at(10),
        ip="198.51.100.9",
        user_agent="Firefox/140",
        meta={"code": "C001", "name": "Acme Group", "changes": {"name": ["Acme", "Acme Group"]}},
    )
    await insert_entry(
        org,
        "payment.recorded",
        actor=owner,
        target_type="payment",
        target_id=uuid.uuid4(),
        created_at=at(20),
        meta={"client_name": "Lovelace Ltd", "amount": 5000},
    )
    return invoice, client_id


async def test_the_csv_export(client):
    owner, org = await practice(client)
    other = await staff(org, "manager", full_name="Grace Hopper", email="grace@example.org")
    _, client_id = await seed(org, owner, other)

    response = await export(client, owner, org, **WINDOW)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    day = today_in(org["timezone"]).isoformat()
    assert response.headers["content-disposition"] == f'attachment; filename="audit-log-{day}.csv"'
    assert response.headers["x-export-row-count"] == "3"
    assert response.headers["x-export-truncated"] == "false"
    assert response.headers["cache-control"] == "no-store"
    # UTF-8 with a byte-order mark and CRLF line ends, for Excel.
    assert response.content.startswith(b"\xef\xbb\xbf")
    text = response.content.decode("utf-8-sig")
    assert text.endswith("\r\n") and "\n" not in text.replace("\r\n", "")

    rows = csv_table(response)
    assert [row["Action"] for row in rows] == [
        "payment.recorded",
        "client.updated",
        "invoice.issued",
    ]
    changed = rows[1]
    assert changed["Time (UTC)"] == "2026-03-10T12:10:00Z"
    assert (changed["User"], changed["Email"]) == ("Grace Hopper", "grace@example.org")
    assert changed["Event"] == "Client details updated"
    assert changed["Summary"] == "Client details updated: C001 · Acme Group"
    assert (changed["Record type"], changed["Record id"]) == ("client", str(client_id))
    assert changed["Changes"] == "name: Acme → Acme Group"
    assert (changed["IP address"], changed["Device"]) == ("198.51.100.9", "Firefox/140")
    assert changed["Request id"].startswith("req-")
    assert json.loads(changed["Details (JSON)"]) == {
        "code": "C001",
        "name": "Acme Group",
        "changes": {"name": ["Acme", "Acme Group"]},
    }
    assert rows[2]["Summary"] == "Invoice issued: INV-00012 (Babbage Ltd)"
    assert rows[2]["Changes"] == ""


async def test_the_xlsx_export(client):
    owner, org = await practice(client)
    other = await staff(org, "manager", full_name="Grace Hopper", email="grace@example.org")
    await seed(org, owner, other)

    response = await export(client, owner, org, "xlsx", **WINDOW)
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    day = today_in(org["timezone"]).isoformat()
    assert response.headers["content-disposition"] == f'attachment; filename="audit-log-{day}.xlsx"'
    assert response.headers["x-export-row-count"] == "3"
    above, rows = xlsx_sheet(response)
    assert above[0][0] == "Audit log"
    notes = [row[0] for row in above[1:] if row[0]]
    assert notes[0] == "Practice: Acme Accountants"
    assert notes[1].startswith("Exported ") and notes[1].endswith(f"UTC by {owner.full_name}")
    assert notes[2].startswith("Filters: From: 2026-03-10T11:59:00+00:00, To: ")
    assert notes[3] == "Entries: 3"
    assert [row["Action"] for row in rows] == [
        "payment.recorded",
        "client.updated",
        "invoice.issued",
    ]
    changed = rows[1]
    assert changed["Time (UTC)"] == "2026-03-10T12:10:00Z"
    assert changed["User"] == "Grace Hopper"
    assert changed["Changes"] == "name: Acme → Acme Group"
    assert changed["IP address"] == "198.51.100.9"


async def test_text_a_spreadsheet_would_run_as_a_formula_is_neutralised(client):
    owner, org = await practice(client)
    sneaky = await create_user("=1+1", "formula@example.com")
    await add_member(org["id"], "staff", user=sneaky)
    await insert_entry(
        org,
        "test.formula",
        actor=sneaky,
        created_at=at(0),
        user_agent="@SUM(1)",
        meta={"note": "x"},
    )
    await insert_entry(
        org, "test.dash", actor=owner, created_at=at(1), user_agent="-2+3", meta={"note": "-1"}
    )

    [dash, formula] = csv_table(await export(client, owner, org, **WINDOW))
    assert formula["User"] == "'=1+1" and formula["Device"] == "'@SUM(1)"
    assert dash["Device"] == "'-2+3"
    _, sheet = xlsx_sheet(await export(client, owner, org, "xlsx", **WINDOW))
    assert sheet[1]["User"] == "'=1+1" and sheet[1]["Device"] == "'@SUM(1)"
    assert sheet[0]["Device"] == "'-2+3"


async def test_the_export_follows_the_list_filters(client):
    owner, org = await practice(client)
    other = await staff(org, "manager", full_name="Grace Hopper", email="grace@example.org")
    invoice, _ = await seed(org, owner, other)

    async def ask(**params):
        return [
            row["Action"]
            for row in csv_table(await export(client, owner, org, **{**WINDOW, **params}))
        ]

    assert await ask() == ["payment.recorded", "client.updated", "invoice.issued"]
    assert await ask(action="client.updated") == ["client.updated"]
    assert await ask(action="invoice.") == ["invoice.issued"]
    assert await ask(actor_user_id=str(other.id)) == ["client.updated"]
    assert await ask(target_type="payment") == ["payment.recorded"]
    assert await ask(target_id=str(invoice)) == ["invoice.issued"]
    assert await ask(q="lovelace") == ["payment.recorded"]
    assert await ask(q="grace") == ["client.updated"]
    assert await ask(**{"from": at(10).isoformat(), "to": at(20).isoformat()}) == [
        "payment.recorded",
        "client.updated",
    ]
    assert await ask(q="nothing like it") == []
    # The workbook says which filters were applied.
    above, _ = xlsx_sheet(
        await export(
            client, owner, org, "xlsx", action="client.", q="grace", **{"from": WINDOW["from"]}
        )
    )
    assert any(
        row[0]
        and row[0].startswith("Filters: Action: client., From: ")
        and row[0].endswith("Search: grace")
        for row in above
    )
    above, _ = xlsx_sheet(
        await export(client, owner, org, "xlsx", actor_user_id=str(other.id), **WINDOW)
    )
    assert any(row[0] and f"User: {other.full_name}" in row[0] for row in above)


async def test_a_large_result_is_cut_to_the_newest_entries(client, monkeypatch):
    owner, org = await practice(client)
    monkeypatch.setattr(audit_export, "MAX_EXPORT_ROWS", 3)
    for index in range(5):
        await insert_entry(org, "test.entry", created_at=at(index), meta={"n": index})

    csv_response = await export(client, owner, org, **WINDOW)
    assert csv_response.headers["x-export-truncated"] == "true"
    assert csv_response.headers["x-export-row-count"] == "3"
    assert [json.loads(row["Details (JSON)"])["n"] for row in csv_table(csv_response)] == [4, 3, 2]

    above, rows = xlsx_sheet(await export(client, owner, org, "xlsx", **WINDOW))
    assert [json.loads(row["Details (JSON)"])["n"] for row in rows] == [4, 3, 2]
    assert any(
        row[0] == "Entries: 3 (the newest 3 of more: narrow the filters to export the rest)"
        for row in above
    )
    [first, _second] = await entry_rows(org, "audit_log.exported")
    assert first["metadata"]["truncated"] is True and first["metadata"]["row_count"] == 3

    # Exactly as many as the cap is not a cut.
    exact = await export(client, owner, org, **{**WINDOW, "from": at(2).isoformat()})
    assert exact.headers["x-export-truncated"] == "false"
    assert len(csv_table(exact)) == 3


async def test_a_huge_detail_is_cut_to_fit_a_cell(client):
    owner, org = await practice(client)
    await insert_entry(org, "test.huge", created_at=at(0), meta={"note": "x" * 40_000})
    [row] = csv_table(await export(client, owner, org, **WINDOW))
    assert len(row["Details (JSON)"]) == audit_export.CELL_LENGTH
    assert row["Details (JSON)"].endswith("…")


async def test_the_export_is_itself_audited(client):
    owner, org = await practice(client)
    await seed(org, owner)
    first = await export(client, owner, org, action="invoice.", **WINDOW)
    assert first.status_code == 200

    [logged] = await entry_rows(org, "audit_log.exported")
    assert logged["actor_user_id"] == str(owner.id)
    assert logged["target_type"] == "audit_log" and logged["target_id"] is None
    assert logged["ip"] == "127.0.0.1" and logged["request_id"]
    assert logged["metadata"] == {
        "format": "csv",
        "row_count": 1,
        "truncated": False,
        "filters": {"action": "invoice.", **WINDOW},
    }
    # The export does not hold its own entry; the next one holds it.
    assert [row["Action"] for row in csv_table(first)] == ["invoice.issued"]
    second = await export(client, owner, org, "xlsx", q="exported")
    _, rows = xlsx_sheet(second)
    assert [row["Event"] for row in rows] == ["Audit log exported"]
    assert json.loads(rows[0]["Details (JSON)"])["row_count"] == 1
    assert [e["metadata"]["format"] for e in await entry_rows(org, "audit_log.exported")] == [
        "csv",
        "xlsx",
    ]


async def test_a_refused_export_is_not_audited(client):
    _, org = await practice(client)
    viewer = await staff(org, "viewer")
    assert (await export(client, viewer, org)).status_code == 403
    assert (await export(client, viewer, org, format="pdf")).status_code in (403, 422)
    assert await entry_rows(org, "audit_log.exported") == []


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
@pytest.mark.parametrize(
    ("role", "status"),
    [
        ("owner", 200),
        ("admin", 200),
        ("partner", 200),
        ("manager", 403),
        ("staff", 403),
        ("billing", 403),
        ("viewer", 403),
    ],
)
async def test_the_export_permission_by_role(client, role, status, fmt):
    owner, org = await practice(client)
    caller = owner if role == "owner" else await staff(org, role)
    assert (await export(client, caller, org, fmt)).status_code == status


async def test_the_export_needs_both_audit_log_permissions(client):
    _, org = await practice(client)
    for name, codes, status in (
        ("Reads", ["audit_log.read"], 403),
        ("Exports", ["audit_log.export"], 403),
        ("Both", ["audit_log.read", "audit_log.export"], 200),
    ):
        await add_custom_role(org["id"], name, codes)
        caller = await staff(org, name)
        assert (await export(client, caller, org)).status_code == status, name
        assert (await export(client, caller, org, "xlsx")).status_code == status, name


async def test_the_export_keeps_practices_apart_and_needs_a_session(client):
    owner, org = await practice(client)
    other_owner, theirs = await practice(client, "Other Practice")
    await insert_entry(org, "test.mine", created_at=at(0), meta={"note": "Babbage"})
    await insert_entry(theirs, "test.theirs", created_at=at(1), meta={"note": "Babbage"})

    assert [
        r["Action"] for r in csv_table(await export(client, owner, org, q="Babbage", **WINDOW))
    ] == ["test.mine"]
    assert [
        r["Action"]
        for r in csv_table(await export(client, other_owner, theirs, q="Babbage", **WINDOW))
    ] == ["test.theirs"]
    assert (await export(client, other_owner, org)).status_code == 404
    assert (await export(client, await create_user(), org)).status_code == 404
    assert (await client.get(audit_url(org, "/export"))).status_code == 401
    # Another practice's export never records into this practice's trail.
    assert len(await entry_rows(org, "audit_log.exported")) == 1


@pytest.mark.parametrize(
    "params",
    [
        {"format": "pdf"},
        {"format": "XLSX"},
        {"format": ""},
        {"from": "yesterday"},
        {"actor_user_id": "x"},
        {"target_id": "x"},
        {"target_type": "Not Valid"},
        {"q": ""},
        {"action": ""},
    ],
)
async def test_the_export_refuses_bad_parameters(client, params):
    owner, org = await practice(client)
    response = await client.get(audit_url(org, "/export"), headers=owner.headers, params=params)
    assert response.status_code == 422, response.text
    assert error(response)["code"] == "VALIDATION_ERROR"
    assert await entry_rows(org, "audit_log.exported") == []


async def test_the_export_leaves_out_cost_figures_for_people_without_rate_read(client):
    owner, org = await practice(client)
    person = await staff(org, full_name="Ada Junior")
    posted = await client.post(
        org_url(org, f"/members/{person.id}/rates"),
        headers=owner.headers,
        json={"effective_from": "2026-01-01", "charge_rate": 10000, "cost_rate": 3500},
    )
    assert posted.status_code == 201, posted.text
    await add_custom_role(org["id"], "Exporter", ["audit_log.read", "audit_log.export"])
    exporter = await staff(org, "Exporter")

    def details(response):
        [row] = [r for r in csv_table(response) if r["Action"] == "staff_rate.created"]
        return json.loads(row["Details (JSON)"])

    full = details(await export(client, owner, org, action="staff_rate.created"))
    assert (full["charge_rate"], full["cost_rate"]) == (10000, 3500)
    cut = details(await export(client, exporter, org, action="staff_rate.created"))
    assert (cut["charge_rate"], cut["cost_rate"]) == (10000, None)
