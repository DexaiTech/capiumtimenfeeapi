"""Excel workbooks: rendering an export and reading an uploaded sheet (no database needed)."""

import io
import zipfile
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services.xlsx_export import (
    XLSX_MEDIA_TYPE,
    XlsxColumn,
    XlsxError,
    read_xlsx,
    render_xlsx,
)

COLUMNS = (
    XlsxColumn("name", "Client"),
    XlsxColumn("minutes", "Hours", "hours"),
    XlsxColumn("value", "Value", "money"),
    XlsxColumn("share_bp", "Share", "bp"),
    XlsxColumn("entries", "Entries", "int"),
    XlsxColumn("on", "Date", "date"),
)


def _sheet(content: bytes):
    return load_workbook(io.BytesIO(content)).worksheets[0]


def test_media_type_is_the_ooxml_spreadsheet_type() -> None:
    assert XLSX_MEDIA_TYPE.endswith("spreadsheetml.sheet")


def test_render_writes_header_rows_and_a_total_row() -> None:
    rows = [
        {
            "name": "Acme",
            "minutes": 90,
            "value": 123456,
            "share_bp": 7500,
            "entries": 3,
            "on": date(2026, 4, 5),
        },
        SimpleNamespace(name="Globex", minutes=30, value=-5, share_bp=None, entries=1, on=None),
    ]
    totals = SimpleNamespace(minutes=120, value=123451, share_bp=5000, entries=4)

    sheet = _sheet(render_xlsx(COLUMNS, rows, totals=totals))

    assert [c.value for c in sheet[1]] == ["Client", "Hours", "Value", "Share", "Entries", "Date"]
    assert sheet["A2"].value == "Acme"
    assert sheet["B2"].value == pytest.approx(1.5)
    assert sheet["C2"].value == pytest.approx(1234.56)
    assert sheet["D2"].value == pytest.approx(75.0)
    assert sheet["E2"].value == 3
    assert sheet["F2"].value.date() == date(2026, 4, 5)
    assert sheet["C3"].value == pytest.approx(-0.05)
    assert sheet["D3"].value is None
    assert sheet["F3"].value is None
    assert [c.value for c in sheet[4]][:5] == ["Total", 2.0, 1234.51, 50.0, 4]
    assert sheet["F4"].value is None


def test_money_hours_and_dates_carry_spreadsheet_formats() -> None:
    sheet = _sheet(
        render_xlsx(COLUMNS, [{"name": "A", "minutes": 60, "value": 100, "on": date(2026, 1, 2)}])
    )
    assert "0.00" in sheet["C2"].number_format
    assert "\u00a3" in sheet["C2"].number_format
    assert sheet["F2"].number_format == "dd/mm/yyyy"
    assert sheet["B2"].number_format == "0.00"


def test_title_and_notes_come_before_the_table_and_the_header_is_frozen() -> None:
    sheet = _sheet(
        render_xlsx(
            COLUMNS, [], sheet="WIP", title="WIP as at 05/04/2026", notes=("Practice: Acme",)
        )
    )
    assert sheet.title == "WIP"
    assert sheet["A1"].value == "WIP as at 05/04/2026"
    assert sheet["A2"].value == "Practice: Acme"
    assert sheet["A4"].value == "Client"
    assert sheet.freeze_panes == "A5"


@pytest.mark.parametrize("text", ["=1+1", "+44 20", "-5", "@SUM(A1)"])
def test_text_that_looks_like_a_formula_is_prefixed(text: str) -> None:
    sheet = _sheet(render_xlsx(COLUMNS, [{"name": text}]))
    assert sheet["A2"].value == f"'{text}"


def test_read_gives_the_header_and_the_rows_without_blank_lines() -> None:
    content = render_xlsx(
        COLUMNS,
        [
            {"name": "  Acme  ", "minutes": 60, "value": 100, "entries": 2},
            {},
            {"name": "Globex", "minutes": 120, "value": 250, "entries": 1},
        ],
    )

    data = read_xlsx(content)

    assert data.headers == ["Client", "Hours", "Value", "Share", "Entries", "Date"]
    assert [row[0] for row in data.rows] == ["Acme", "Globex"]  # the blank line is dropped
    assert data.rows[0][1] == pytest.approx(1.0)


def test_read_skips_leading_blank_rows_and_pads_short_rows() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append([None, None])
    sheet.append(["Code", "Name", "Fee"])
    sheet.append(["C1", "Acme"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    data = read_xlsx(buffer.getvalue())

    assert data.headers == ["Code", "Name", "Fee"]
    assert data.rows == [["C1", "Acme", None]]


def test_read_refuses_a_file_that_is_not_a_workbook() -> None:
    with pytest.raises(XlsxError):
        read_xlsx(b"not a workbook")


def test_read_refuses_too_many_rows() -> None:
    content = render_xlsx(COLUMNS, [{"name": f"C{i}"} for i in range(5)])
    with pytest.raises(XlsxError, match="more than 3 rows"):
        read_xlsx(content, max_rows=3)


def test_read_refuses_an_archive_that_unpacks_to_an_unreasonable_size() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/big.bin", b"\0" * (61 * 1024 * 1024))
    with pytest.raises(XlsxError, match="too large"):
        read_xlsx(buffer.getvalue())


def test_money_is_exact_for_large_pence_values() -> None:
    sheet = _sheet(render_xlsx(COLUMNS, [{"value": 123456789012}]))
    assert Decimal(str(sheet["C2"].value)) == Decimal("1234567890.12")
