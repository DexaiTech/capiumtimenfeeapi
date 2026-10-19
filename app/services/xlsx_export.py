"""Excel (.xlsx) workbooks: exports, import templates and reading uploaded sheets.

Numbers stay numbers so a spreadsheet can add them up: money is written as pounds
(``123456`` -> ``1234.56``) with a pound format, minutes as decimal hours, basis points as a
percentage, dates as real dates shown ``dd/mm/yyyy``. Text that a spreadsheet would read as a
formula is prefixed with an apostrophe, as in the CSV exports.
"""

import io
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

Kind = Literal["text", "money", "hours", "minutes", "int", "bp", "date"]

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_FORMATS: dict[str, str] = {
    "money": '"\u00a3"#,##0.00;[Red]-"\u00a3"#,##0.00',
    "hours": "0.00",
    "minutes": "0",
    "int": "0",
    "bp": '0.00"%"',
    "date": "dd/mm/yyyy",
}
_MAX_WIDTH = 60
_MAX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
_HEADER_FILL = PatternFill("solid", start_color="EDE9F7", end_color="EDE9F7")


class XlsxError(ValueError):
    """The file is not a readable workbook, or it is too large to read."""


@dataclass(frozen=True, slots=True)
class XlsxColumn:
    key: str
    header: str
    kind: Kind = "text"
    width: int | None = None


@dataclass(frozen=True, slots=True)
class SheetData:
    """The first sheet of a workbook: its header row and the rows below it (blank rows dropped)."""

    headers: list[str]
    rows: list[list[Any]]


def _get(row: Mapping[str, Any] | object, key: str) -> Any:
    if isinstance(row, Mapping):
        return row.get(key)
    return getattr(row, key, None)


def _text(value: Any) -> str:
    text = str(value)
    return f"'{text}" if text.startswith(_FORMULA_START) else text


def _number(value: Any, kind: Kind) -> Any:
    if value is None or value == "":
        return None
    match kind:
        case "money":
            return Decimal(int(value)) / 100
        case "hours":
            return (Decimal(int(value)) / 60).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        case "bp":
            return Decimal(int(value)) / 100
        case "minutes" | "int":
            return int(value)
        case "date":
            return value if isinstance(value, date | datetime) else None
        case _:
            return _text(value)


def render_xlsx(
    columns: Sequence[XlsxColumn],
    rows: Iterable[Mapping[str, Any] | object],
    *,
    sheet: str = "Report",
    title: str | None = None,
    notes: Sequence[str] = (),
    totals: Mapping[str, Any] | object | None = None,
) -> bytes:
    """One sheet: an optional title and notes, a bold header row, the rows, a final Total row.

    The header row is frozen and filtered. ``totals`` supplies the numeric columns of the last
    row (text and date columns stay empty; the first column reads "Total").
    """
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:  # pragma: no cover - a new workbook always has an active sheet
        raise XlsxError("The workbook has no sheet.")
    worksheet.title = sheet[:31] or "Sheet1"
    workbook.properties.creator = "Capium Time & Fees"

    line = 1
    if title:
        worksheet.cell(row=line, column=1, value=_text(title)).font = Font(bold=True, size=14)
        line += 1
    for note in notes:
        worksheet.cell(row=line, column=1, value=_text(note)).font = Font(italic=True)
        line += 1
    if title or notes:
        line += 1  # a blank row between the notes and the table

    header_row = line
    widths = [len(column.header) for column in columns]
    for index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=header_row, column=index, value=column.header)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(
            horizontal="left" if column.kind == "text" else "right", vertical="center"
        )

    last_row = header_row
    for row in rows:
        last_row += 1
        for index, column in enumerate(columns, start=1):
            value = _number(_get(row, column.key), column.kind)
            cell = worksheet.cell(row=last_row, column=index, value=value)
            if column.kind in _FORMATS:
                cell.number_format = _FORMATS[column.kind]
            if value is not None:
                widths[index - 1] = max(widths[index - 1], len(str(value)))

    if totals is not None:
        last_row += 1
        top = Border(top=Side(style="thin"))
        for index, column in enumerate(columns, start=1):
            total: Any
            if index == 1:
                total = "Total"
            elif column.kind in ("text", "date"):
                total = None
            else:
                total = _number(_get(totals, column.key), column.kind)
            cell = worksheet.cell(row=last_row, column=index, value=total)
            cell.font = Font(bold=True)
            cell.border = top
            if column.kind in _FORMATS:
                cell.number_format = _FORMATS[column.kind]

    for index, column in enumerate(columns, start=1):
        width = column.width or min(widths[index - 1] + 2, _MAX_WIDTH)
        worksheet.column_dimensions[get_column_letter(index)].width = max(width, 8)
    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
    if columns:
        worksheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(columns))}{max(last_row, header_row)}"
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_xlsx(content: bytes, *, max_rows: int = 5000, max_columns: int = 60) -> SheetData:
    """The header row and the rows of the first sheet of an uploaded workbook.

    Raises ``XlsxError`` for a file that is not a workbook, one that unpacks to an unreasonable
    size, or one with more than ``max_rows`` data rows. Cells hold Python values (text, numbers,
    dates, None); header cells are trimmed text; rows that are entirely blank are dropped.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > _MAX_UNCOMPRESSED_BYTES:
                raise XlsxError("The workbook is too large to read.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except XlsxError:
        raise
    except Exception as exc:  # openpyxl raises many types for a damaged or foreign file
        raise XlsxError("The file is not a valid Excel workbook.") from exc

    try:
        worksheet = workbook.worksheets[0] if workbook.worksheets else None
        if worksheet is None:
            raise XlsxError("The workbook has no sheet.")
        headers: list[str] = []
        rows: list[list[Any]] = []
        for raw in worksheet.iter_rows(max_col=max_columns, values_only=True):
            cells = list(raw)
            if not headers:
                if all(cell is None or str(cell).strip() == "" for cell in cells):
                    continue
                headers = ["" if cell is None else str(cell).strip() for cell in cells]
                while headers and headers[-1] == "":
                    headers.pop()
                continue
            if all(
                cell is None or (isinstance(cell, str) and cell.strip() == "") for cell in cells
            ):
                continue
            if len(rows) >= max_rows:
                raise XlsxError(f"The sheet has more than {max_rows} rows.")
            width = len(headers)
            row = [cell.strip() if isinstance(cell, str) else cell for cell in cells[:width]]
            row.extend([None] * (width - len(row)))
            rows.append(row)
        return SheetData(headers=headers, rows=rows)
    finally:
        workbook.close()
